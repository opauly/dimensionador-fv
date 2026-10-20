"""Regenerate calculations/nec_tables.py from Oscar's "Dimensionador CCC + FV" workbook.

    python3 scripts/gen_nec_tables.py "<path to Dimensionador CCC + FV v0.2.xlsx>"

The NEC / IEEE tables in the app are a verbatim digest of the workbook's own table
sheets (Cu only — the app assumes copper conductors and PVC conduit), so the
wiring numbers stay tied to the spreadsheet. Run this again after the workbook's
tables are corrected; never hand-edit the generated module.
"""
from __future__ import annotations

import pprint
import sys
from datetime import date
from pathlib import Path

import openpyxl

DEFAULT_SRC = (
    "~/Desktop/Pauly & Co/Clientes/María Luisa Montero/María Alexandra Montero/"
    "Dimensionador CCC + FV v0.2 2.xlsx"
)
OUT = Path(__file__).resolve().parent.parent / "calculations" / "nec_tables.py"


def _g(v) -> str:
    """Gauge cells are strings in the workbook ('1/0', '250'); normalise floats like 6.0."""
    s = str(v).strip()
    return s[:-2] if s.endswith(".0") else s


def _num(v):
    return None if v in (None, "", "-") else float(v)


def main(src: str) -> None:
    wb = openpyxl.load_workbook(Path(src).expanduser(), data_only=True)

    # T310.15(B)(16) — Cu: (gauge, 60C, 75C, 90C), largest -> smallest as in the sheet
    ws = wb["T310.15(B)(16)"]
    t310 = []
    for r in range(4, 34):
        t310.append((_g(ws.cell(r, 2).value), _num(ws.cell(r, 1).value), _num(ws.cell(r, 3).value), _num(ws.cell(r, 5).value)))
    # T8 — gauge -> circular mils (A:B rows 3-32) and DC resistance Cu ohm/km (L:M rows 4-33)
    ws = wb["T8"]
    cmils = {_g(ws.cell(r, 1).value): float(ws.cell(r, 2).value) for r in range(3, 33)}
    r_dc = {_g(ws.cell(r, 12).value): float(ws.cell(r, 13).value) for r in range(4, 34)}
    # T9 — AC: reactance (PVC/Al conduit) col B, resistance Cu in PVC col D, rows 5-25
    ws = wb["T9"]
    t9 = {_g(ws.cell(r, 1).value): (float(ws.cell(r, 2).value), float(ws.cell(r, 4).value)) for r in range(5, 26)}
    # T250.66 Cu — grounding electrode conductor by largest ungrounded conductor (rows 3-32)
    ws = wb["T250.66"]
    t25066 = {_g(ws.cell(r, 1).value): _g(ws.cell(r, 2).value) for r in range(3, 33)}
    # T250.122 Cu — equipment grounding conductor by OCPD (rows 3-21)
    ws = wb["T250.122"]
    t250122 = [(int(ws.cell(r, 1).value), _g(ws.cell(r, 2).value)) for r in range(3, 22)]
    # Standard OCPD sizes (Breakers sheet col B, ascending)
    ws = wb["Breakers"]
    breakers = sorted({int(ws.cell(r, 2).value) for r in range(2, 39)})
    # T5 conductor areas mm2 (only the insulations the FV sheets use)
    areas = {}
    for name in ("RHW-2", "THWN-2"):
        ws = wb[name]
        areas[name] = {_g(ws.cell(r, 1).value): _num(ws.cell(r, 2).value) for r in range(3, 33)}
    # T4 PVC conduit usable area (40 % fill) mm2 by trade size
    conduits = {}
    for name in ("PVC A", "PVC 40", "PVC 80"):
        ws = wb[name]
        rows = []
        for r in range(3, 16):
            area = _num(ws.cell(r, 1).value)
            if area is not None:
                rows.append((_g(ws.cell(r, 3).value), _g(ws.cell(r, 2).value), area))
        rows.sort(key=lambda t: t[2])
        conduits[name] = rows

    order = ["18", "16", "14", "12", "10", "8", "6", "4", "3", "2", "1", "1/0", "2/0", "3/0", "4/0",
             "250", "300", "350", "400", "500", "600", "700", "750", "800", "900", "1000", "1250", "1500", "1750", "2000"]
    by_gauge = {g: (a60, a75, a90) for g, a60, a75, a90 in t310}
    t310_ordered = [(g, *by_gauge[g]) for g in order]

    def lit(name, value):
        return f"{name} = " + pprint.pformat(value, width=110, sort_dicts=False) + "\n\n"

    body = (
        '"""NEC / IEEE tables for the wiring summary — GENERATED, do not hand-edit.\n\n'
        "Source: Oscar's \"Dimensionador CCC + FV\" workbook (v0.2), digested by\n"
        f"scripts/gen_nec_tables.py on {date.today().isoformat()}. Copper conductors and PVC conduit only.\n"
        "Table numbers follow the workbook (NEC 2014/2017 numbering, e.g. T310.15(B)(16)).\n"
        '"""\n\n'
        "# Gauges smallest -> largest; every per-gauge table below is keyed by these strings.\n"
        + lit("GAUGE_ORDER", order)
        + "# T310.15(B)(16), Cu: gauge -> (ampacity 60C, 75C, 90C); None = not listed for that column.\n"
        + lit("T310_15_B_16_CU", {g: (a60, a75, a90) for g, a60, a75, a90 in t310_ordered})
        + "# T8: gauge -> area in circular mils; and DC resistance of uncoated Cu in ohm/km.\n"
        + lit("T8_CMILS", {g: cmils[g] for g in order})
        + lit("T8_R_DC_CU_OHM_KM", {g: r_dc[g] for g in order if g in r_dc})
        + "# T9 (AC): gauge -> (reactance ohm/km in PVC/Al conduit, resistance Cu ohm/km in PVC conduit).\n"
        + lit("T9_AC_PVC_CU", {g: t9[g] for g in order if g in t9})
        + "# T250.66, Cu: largest ungrounded conductor gauge -> grounding conductor gauge.\n"
        + lit("T250_66_CU", {g: t25066[g] for g in order})
        + "# T250.122, Cu: (OCPD rating A, EGC gauge), largest first as in the workbook.\n"
        + lit("T250_122_CU", t250122)
        + "# Standard OCPD ratings (A), ascending (workbook 'Breakers' sheet).\n"
        + lit("OCPD_SIZES_A", breakers)
        + "# T5: conductor cross-section in mm2 by insulation type (None = not listed).\n"
        + lit("T5_AREA_MM2", {k: {g: v[g] for g in order} for k, v in areas.items()})
        + "# T4: PVC conduit usable fill area (mm2) -> (trade size in, metric designator), smallest first.\n"
        + lit("T4_PVC_CONDUIT", conduits)
    )
    OUT.write_text(body)
    print("wrote", OUT, len(body), "bytes")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SRC)
